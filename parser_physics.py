import sys
import os
import re
import json
from datetime import datetime, timezone
import openpyxl

from base_parser import (
    HEADERS,
    DAYS_ORDER,
    get_latest_xlsx_url as base_get_latest_xlsx_url,
    download_file,
    compute_semester_string,
    save_database_with_cache_check
)

PHYSICS_URL = "https://kpfu.ru/physics/raspisanie-zanyatij"
OUTPUT_FILE = "schedule_physics.json"

def get_latest_xlsx_url(retries=3):
    return base_get_latest_xlsx_url(PHYSICS_URL, HEADERS, retries=retries)

TEACHER_RE = re.compile(
    r'\b([А-ЯЁа-яё][а-яё]+(?:-[А-ЯЁа-яё][а-яё]+)?)\s*([А-ЯЁ])\.\s*([А-ЯЁ])\.?(?=[^а-яА-Яa-zA-Z]|$)|'
    r'\b([А-ЯЁ])\.\s*([А-ЯЁ])\.?\s+([А-ЯЁа-яё][а-яё]+(?:-[А-ЯЁа-яё][а-яё]+)?)|'
    r'\b([А-ЯЁа-яё][а-яё]+(?:-[А-ЯЁа-яё][а-яё]+)?)\s*([А-ЯЁ])\.(?![а-яА-Яa-zA-Z])'
)

TEACHER_ALIASES = {
    'Алакин А.Б.': 'Балакин А.Б.',
    'Воронина': 'Воронина Е.В.'
}

BUILDING_MAP = [
    (re.compile(r'карла\s+маркса[\s,]+74[а-яА-Я]?|сц\b', re.I), "Спортивный центр (Карла Маркса, 74а)"),
    (re.compile(r'кремлевск\w*[\s,]+35|2[\s-]*й\s+корпус', re.I), "2-й учебный корпус (Кремлевская, 35)"),
    (re.compile(r'лево[- ]булачн\w*[\s,]+44', re.I), "Лево-Булачная, 44"),
    (re.compile(r'межлаук\w*[\s,]+1', re.I), "Межлаука, 1"),
    (re.compile(r'кск\s+уникс\b|уникс\b', re.I), "КСК УНИКС"),
    (re.compile(r'химическ\w*\s+институт\w*|хим\.?\s*ин-т', re.I), "Химический институт"),
    (re.compile(r'гл\.?\s*з(?:д)?\.?|главн\w*\s+здан\w*', re.I), "Главное здание"),
    (re.compile(r'\d*\s*конф\w*[- ]зал\b|конф\w*[- ]зал\b|к\.?-?з\.?', re.I), "Конференц-зал"),
    (re.compile(r'лицей\b', re.I), "Лицей им. Лобачевского"),
    (re.compile(r'ляф\b', re.I), "Лаборатория ядерной физики (ЛЯФ)"),
    (re.compile(r'фиц\s+каз\s+нц\s+ран|каз\s+нц\s+ран', re.I), "ФИЦ КазНЦ РАН"),
]

def extract_teachers_with_raw(text):
    teachers = []
    raws = []
    for m in TEACHER_RE.finditer(text):
        raw_match = m.group(0).strip()
        if m.group(1):
            s_name = m.group(1).capitalize()
            t_name = f"{s_name} {m.group(2)}.{m.group(3)}."
        elif m.group(4):
            s_name = m.group(6).capitalize()
            t_name = f"{s_name} {m.group(4)}.{m.group(5)}."
        elif m.group(7):
            s_name = m.group(7).capitalize()
            t_name = f"{s_name} {m.group(8)}."
        else:
            continue
        t_name = TEACHER_ALIASES.get(t_name, t_name)
        teachers.append(t_name)
        raws.append(raw_match)

    # Check for Voronina without initials
    m_vor = re.search(r'\bВоронина\b(?!\s+[А-ЯЁ]\.)', text)
    if m_vor and 'Воронина Е.В.' not in teachers:
        teachers.append('Воронина Е.В.')
        raws.append(m_vor.group(0))

    res_t, res_r = [], []
    for t, r in zip(teachers, raws):
        if t not in res_t:
            res_t.append(t)
            res_r.append(r)
    return res_t, res_r

def extract_building(text):
    for pat, name in BUILDING_MAP:
        if pat.search(text):
            return name
    return "Кремлевская, 16А (Институт физики)"

def extract_room(text):
    cleaned = re.sub(r'(?:с\s+)?\b\d{1,2}(?:\s*-\s*\d{1,2})?(?:\s*,\s*\d{1,2}(?:\s*-\s*\d{1,2})?)*\s*н\.?', ' ', text, flags=re.I)
    
    # 1. Explicit aud
    m = re.search(r'ауд\.?\s*([0-9А-Яа-яA-Za-z]+(?:\s*,\s*[0-9А-Яа-яA-Za-z]+)*)', cleaned, re.I)
    if m:
        return m.group(1).strip()

    # 2. 1 конф.-зал / 1 конференц-зал
    m_kz = re.search(r'\b(\d{1,2})\s*конф[\w.]*[- ]?зал\b', cleaned, re.I)
    if m_kz:
        return f"{m_kz.group(1)} конф.-зал"

    # 3. ЛЯФ rooms
    m_lyaf = re.search(r'ляф[\s,]*(?:к\.?|комн?\.?)?\s*([0-9\s,]+)', cleaned, re.I)
    lyaf_room = ""
    if m_lyaf and m_lyaf.group(1).strip():
        nums = [n.strip() for n in m_lyaf.group(1).split(',') if n.strip()]
        lyaf_room = "к. " + ", ".join(nums)

    # 4. Multi-digit rooms or single digit + letter
    m_multi = re.findall(r'\b([0-9]{3,4}[а-мА-Мо-яО-Яa-zA-Z]?|[1-9][а-мА-Мо-яО-Яa-zA-Z])\b', cleaned)
    valid_rooms = [r for r in m_multi if r not in ["2024", "2025", "2026", "2027", "2028"]]
    if valid_rooms:
        res = ", ".join(valid_rooms)
        if lyaf_room:
            return f"{res}, {lyaf_room}"
        return res
    if lyaf_room:
        return lyaf_room
    return ""

def extract_weeks(text):
    week_start = 1
    week_end = 18
    week_type = "all"
    low = text.lower()
    
    if "ч.н" in low or "ч/н" in low or "ч\\н" in low or "четн" in low:
        week_type = "even"
    elif "н.н" in low or "н/н" in low or "н\\н" in low or "нечет" in low:
        week_type = "odd"

    all_nums = []
    for m in re.finditer(r'(?:с\s+)?(\d{1,2})\s*-\s*(\d{1,2})', text, re.I):
        after = text[m.end():m.end()+15].lower()
        if 'н' in after or 'нед' in after or 'пн' in after or 'вт' in after:
            all_nums.extend([int(m.group(1)), int(m.group(2))])
            
    for m in re.finditer(r'(\d{1,2})\s*(?:н\b|нед|н\.)', text, re.I):
        val = int(m.group(1))
        prev_str = text[max(0, m.start()-15):m.start()]
        comma_nums = re.findall(r'\b(\d{1,2})\s*,', prev_str)
        for cn in comma_nums:
            all_nums.append(int(cn))
        all_nums.append(val)

    s_match = re.search(r'\bс\s+(\d{1,2})\s*н\b', text, re.I)
    if s_match:
        all_nums.extend([int(s_match.group(1)), 18])

    if all_nums:
        week_start = min(all_nums)
        week_end = max(all_nums)
        if week_type == "all":
            if week_start == 1 and week_end == 9:
                week_type = "first_half"
            elif week_start == 10 and week_end == 18:
                week_type = "second_half"

    return week_start, week_end, week_type

def extract_subgroup(text):
    m = re.search(r'\b([12])\s*(?:гр|подгр|п/г)', text, re.I)
    if m:
        return int(m.group(1))
    return None

def extract_lesson_type(text):
    low = text.lower()
    if "(лаб" in low or "лаб." in low or "лаборат" in low or re.search(r'\b-\s*лаб\b', low):
        return "lab"
    elif "(пр" in low or "пр." in low or "практ" in low or re.search(r'\b-\s*пр\b', low):
        return "practice"
    elif "(л)" in low or "(л+" in low or "лек." in low or "лекция" in low or re.search(r'\b-\s*(?:лек|л\b)', low):
        return "lecture"
    elif "дистанцион" in low or "онлайн" in low:
        return "distance"
    elif "эор" in low:
        return "eor"
    elif "цор" in low:
        return "cor"
    elif "физическая культура" in low or "спорт" in low:
        return "practice"
    return "other"

def clean_subject(raw_text, raw_teachers, room):
    s = raw_text
    # 1. Strip elective header prefix
    s = re.sub(r'К[у]?рс\s+по\s+выбору\s*(?:\([^)]*\))?:?', ' ', s, flags=re.IGNORECASE)

    # 2. Strip teachers
    for rt in raw_teachers:
        s = s.replace(rt, ' ')
        parts = rt.split()
        if parts:
            s = re.sub(r'\b' + re.escape(parts[0]) + r'\b', ' ', s)
    s = re.sub(r'\b[А-ЯЁа-яё][а-яё]+(?:-[А-ЯЁа-яё][а-яё]+)?\s*[А-ЯЁ]\.\s*[А-ЯЁ]\.?', ' ', s)
    s = re.sub(r'\b[А-ЯЁ]\.\s*[А-ЯЁ]\.?\s+[А-ЯЁа-яё]+', ' ', s)
    s = re.sub(r'\b[А-ЯЁа-яё][а-яё]+\s*[А-ЯЁ]\.?', ' ', s)
    s = re.sub(r'\b[А-ЯЁ]\.[А-ЯЁ]\.?\b', ' ', s)
    s = re.sub(r'\bВоронина\b', ' ', s)
    
    # 3. Strip URLs, times, weeks
    s = re.sub(r'https?://\S+', ' ', s)
    s = re.sub(r'\b\d{1,2}[.:]\d{2}\s*-\s*\d{1,2}[.:]\d{2}\b', ' ', s)
    s = re.sub(r'(?:с\s+)?\b\d{1,2}(?:\s*-\s*\d{1,2})?(?:\s*,\s*\d{1,2}(?:\s*-\s*\d{1,2})?)*\s*н\.?', ' ', s, flags=re.I)
    s = re.sub(r'\bс\s+\d{1,2}\s*н\.?', ' ', s, flags=re.I)
    
    # 4. Strip rooms & buildings
    if room:
        for r_part in room.split(','):
            r_part = r_part.strip()
            if r_part:
                s = re.sub(r'\b' + re.escape(r_part) + r'\b', ' ', s)
    s = re.sub(r'\bауд\.?\s*', ' ', s, flags=re.I)
    s = re.sub(r'\bк\.\s*\d+\b', ' ', s, flags=re.I)
    s = re.sub(r'\bкомн?\.?\s*\d+\b', ' ', s, flags=re.I)
    s = re.sub(r'\b[0-9]{3,4}[а-яА-Яa-zA-Z]?\b', ' ', s)
    s = re.sub(r'\b[1-9][а-яА-Яa-zA-Z]\b', ' ', s)
    s = re.sub(r'\((?:л|пр|лаб|л\+пр|лек|практ|дист|онлайн)\)', ' ', s, flags=re.I)
    s = re.sub(r'\b-\s*(?:л|пр|лаб|лек|практ)\.?', ' ', s, flags=re.I)
    
    for b_kw in [
        r'карла\s+маркса[\s,]+74[а-яА-Я]?', r'сц\b',
        r'кремлевск\w*[\s,]+35', r'кремлевская[\s,]+16[а-яА-Я]?',
        r'лево[- ]булачн\w*[\s,]+44', r'межлаук\w*[\s,]+1',
        r'кск\s+уникс\b', r'уникс\b', r'химическ\w*\s+институт\w*',
        r'хим\.?\s*ин-т', r'гл\.?\s*з(?:д)?\.?', r'главн\w*\s+здан\w*',
        r'\d*\s*конф[\w.]*[- ]?зал\b', r'к\.?-?з\.?', r'лицей\b',
        r'ляф[\s,]*(?:к\.?|комн?\.?)?\s*[0-9\s,]*',
        r'фиц\s+каз\s+нц\s+ран|каз\s+нц\s+ран'
    ]:
        s = re.sub(b_kw, ' ', s, flags=re.I)
        
    s = re.sub(r'\b\d/\d\*?\s*гр\.?', ' ', s, flags=re.I)
    s = re.sub(r'\b\d\s*(?:подгр|гр|п/г)\.?', ' ', s, flags=re.I)
    s = re.sub(r'\b[чн]\.[н\.]*', ' ', s, flags=re.I)
    s = re.sub(r'\b[чн]/[н\.]*', ' ', s, flags=re.I)
    s = re.sub(r'\b[чн]\\[н\.]*', ' ', s, flags=re.I)
    s = re.sub(r'[,;\n\r\t:]+', ' ', s)
    s = re.sub(r'\s+', ' ', s).strip(' ,.-/()')
    return s

def is_pure_elective_header(l):
    low = l.strip().lower()
    return bool(re.match(r'^(?:к[у]?рс|дисциплина)\s+по\s+выбору\s*:?\s*(?:\([^)]*\))?\s*:?$', low))

def is_building_or_room_line(l):
    for pat, _ in BUILDING_MAP:
        if pat.search(l):
            return True
    if re.search(r'конф\b', l, re.I):
        return True
    if re.match(r'^\s*(?:[cс]\s+)?\d+(?:[\s,-]+\d+)*\s*н\b', l, re.I):
        if not re.search(r'\b[А-ЯЁ][а-яё]{3,}.*?\((?:л|пр|лаб|лек|практ)\)', l):
            return True
    return False

def is_subject_header_line(l):
    low = l.strip().lower()
    if is_building_or_room_line(l):
        return False
    if is_pure_elective_header(l):
        return False
    if re.match(r'^\((?:л|пр|лаб|л\+пр|лек|практ)\)', low):
        return False
    if re.search(r'[А-ЯЁа-яё]{3,}.*?\((?:л|пр|лаб|л\+пр|лек|практ)\)', l):
        return True
    if any(low.startswith(p) for p in ["курс по выбору", "крс по выбору", "дисциплина по выбору", "физическая культура", "элективные курсы"]):
        return True
    if re.match(r'^\d{1,2}(?:-\d{1,2})?\s*н\.?\s+[А-ЯЁ]', l):
        return True
    t_list, _ = extract_teachers_with_raw(l)
    if not t_list and not re.search(r'\b\d{1,2}-\d{1,2}н\b', l) and not re.search(r'\b[0-9]{3,4}\b', l):
        if len(re.findall(r'[А-ЯЁа-яё]{3,}', l)) >= 1:
            return True
    return False

def split_semicolons_outside_parens(text):
    res = []
    in_paren = 0
    for char in text:
        if char == '(':
            in_paren += 1
            res.append(char)
        elif char == ')':
            if in_paren > 0:
                in_paren -= 1
            res.append(char)
        elif char == ';' and in_paren == 0:
            res.append('\n')
        else:
            res.append(char)
    return ''.join(res)

def expand_colon_line(line):
    # Pattern: [weeks] [Subject]: [sub-lessons separated by ;]
    m = re.match(r'^(?:(?:\d{1,2}(?:-\d{1,2})?\s*н\.?\s+)?)([^:]+):\s*(.+)$', line)
    if not m:
        return [line]
    base_subj = m.group(1).strip()
    rest = m.group(2).strip()
    
    if not re.search(r'(?:с\s+)?\d+.*(?:н|нед).*-(?:\s*лек|\s*л\b|\s*пр|\s*лаб|\s*вся)', rest, re.I):
        return [line]
        
    parts = [p.strip() for p in rest.split(';') if p.strip()]
    trailing_room = ''
    if parts and (parts[-1].startswith('ауд') or re.match(r'^[0-9]{3,4}[а-яА-Яa-zA-Z]?$', parts[-1]) or re.match(r'^[1-9][а-яА-Яa-zA-Z]$', parts[-1])):
        trailing_room = parts.pop()

    expanded = []
    for p in parts:
        if re.search(r'\b[А-ЯЁ][а-яё]{4,}', p) and not re.search(r'-\s*(?:лек|л\b|пр|лаб)', p):
            expanded.append(f'{p} {trailing_room}')
        else:
            expanded.append(f'{base_subj} {p} {trailing_room}')
    return expanded

def parse_physics_cell(cell_text):
    trimmed = cell_text.strip()
    if not trimmed:
        return []
    
    lower = trimmed.lower()
    if "____" in trimmed or "департамент образования" in lower or "турилова" in lower or "поминов" in lower:
        return []

    # 1. Expand colons if present
    expanded_colon_lines = []
    for raw_l in cell_text.split('\n'):
        raw_l = raw_l.strip()
        if not raw_l: continue
        expanded_colon_lines.extend(expand_colon_line(raw_l))
    
    text_colon_expanded = '\n'.join(expanded_colon_lines)

    # 2. Split semicolons outside parentheses into newlines
    text_semi = split_semicolons_outside_parens(text_colon_expanded)

    # 3. Split in-line subjects pasted without newlines (e.g. "... 809 Спутниковые системы... (пр)")
    # Insert \n before a subject header pattern only when preceded by room, teacher initial, or building
    text_separated = re.sub(
        r'(\b[0-9]{3,4}[а-яА-Яa-zA-Z]?|[А-ЯЁ]\.|гл\.?\s*з(?:д)?\.?|уникс|зал)\s+([А-ЯЁ][а-яё]{3,}.*?\((?:л|пр|лаб|л\+пр|лек|практ)\))',
        r'\1\n\2',
        text_semi
    )
    text_separated = re.sub(
        r'(\b[0-9]{3,4}[а-яА-Яa-zA-Z]?|[А-ЯЁ]\.|гл\.?\s*з(?:д)?\.?|уникс|зал)\s+(\d{1,2}(?:-\d{1,2})?\s*н\.?\s+[А-ЯЁ][а-яё]{3,})',
        r'\1\n\2',
        text_separated
    )
    text_separated = re.sub(
        r'(\b[0-9]{3,4}[а-яА-Яa-zA-Z]?|[А-ЯЁ]\.|гл\.?\s*з(?:д)?\.?|уникс|зал)\s+(К[у]?рс\s+по\s+выбору\s*(?:\([^)]*\))?:?\s*[А-ЯЁ])',
        r'\1\n\2',
        text_separated
    )

    raw_lines = [l.strip() for l in text_separated.split('\n') if l.strip()]
    if not raw_lines:
        return []

    # Join lowercase continuation lines and extract elective category
    lines = []
    elective_prefix = ""
    for l in raw_lines:
        if is_pure_elective_header(l):
            elective_prefix = "Курс по выбору: "
            continue
        if lines and (l[0].islower() and not l.startswith('ауд')) and not l.startswith('(л') and not l.startswith('(пр') and not l.startswith('(лаб'):
            lines[-1] = lines[-1] + ' ' + l
        else:
            lines.append(l)

    sections = []
    cur_sec = []
    for l in lines:
        if is_subject_header_line(l) and cur_sec:
            sections.append(cur_sec)
            cur_sec = [l]
        else:
            cur_sec.append(l)
    if cur_sec:
        sections.append(cur_sec)

    results = []
    for sec in sections:
        sec_text = '\n'.join(sec)
        first_line = sec[0]
        url_match = re.search(r'(https?://[^\s)"]+)', sec_text)
        url = url_match.group(1).rstrip('.,;') if url_match else ""
        is_additional = bool(re.search(r"\s*-\s*д\.?(?:\s|$)", sec_text.lower()))

        type_lines = [(i, l) for i, l in enumerate(sec) if re.match(r'^\((?:л|пр|лаб|л\+пр|лек|практ)\)', l.strip().lower())]
        teacher_lines = []
        for i, l in enumerate(sec):
            t_list, r_list = extract_teachers_with_raw(l)
            if t_list:
                teacher_lines.append((i, l, t_list, r_list))

        common_teachers, common_raw_t = extract_teachers_with_raw(sec_text)
        base_subj = clean_subject(first_line, common_raw_t, '')
        if elective_prefix and not base_subj.lower().startswith("курс по выбору"):
            base_subj = elective_prefix + base_subj

        if len(type_lines) >= 2:
            t_names = ', '.join(common_teachers)
            for _, tl in type_lines:
                w_start, w_end, w_type = extract_weeks(tl)
                results.append({
                    'subject': base_subj,
                    'rawText': f'{base_subj}\n{tl}',
                    'teacher': t_names,
                    'room': extract_room(tl) or extract_room(sec_text),
                    'building': extract_building(tl) or extract_building(sec_text),
                    'type': extract_lesson_type(tl),
                    'url': url,
                    'isAdditional': is_additional,
                    'weekStart': w_start,
                    'weekEnd': w_end,
                    'weekType': w_type,
                    'subgroup': extract_subgroup(tl)
                })
        elif len(teacher_lines) >= 2 and any(extract_subgroup(l) or extract_room(l) for _, l, _, _ in teacher_lines):
            subj_type = extract_lesson_type(sec_text)
            for idx_t, (line_idx, l_str, t_list, r_list) in enumerate(teacher_lines):
                next_idx = teacher_lines[idx_t+1][0] if idx_t+1 < len(teacher_lines) else len(sec)
                sub_lines = sec[line_idx:next_idx]
                sub_text = '\n'.join(sub_lines)
                w_start, w_end, w_type = extract_weeks(sub_text)
                results.append({
                    'subject': base_subj,
                    'rawText': f'{base_subj}\n{sub_text}',
                    'teacher': ', '.join(t_list),
                    'room': extract_room(sub_text) or extract_room(sec_text),
                    'building': extract_building(sub_text) or extract_building(sec_text),
                    'type': subj_type,
                    'url': url,
                    'isAdditional': is_additional,
                    'weekStart': w_start,
                    'weekEnd': w_end,
                    'weekType': w_type,
                    'subgroup': extract_subgroup(sub_text)
                })
        else:
            room = extract_room(sec_text)
            bld = extract_building(sec_text)
            w_start, w_end, w_type = extract_weeks(sec_text)
            subgrp = extract_subgroup(sec_text)
            l_type = extract_lesson_type(sec_text)
            clean_s = clean_subject(first_line, common_raw_t, room)
            if len(clean_s) < 2:
                clean_s = clean_subject(sec_text, common_raw_t, room)
            if elective_prefix and not clean_s.lower().startswith("курс по выбору"):
                clean_s = elective_prefix + clean_s
            if len(clean_s) >= 2:
                results.append({
                    'subject': clean_s,
                    'rawText': sec_text,
                    'teacher': ', '.join(common_teachers),
                    'room': room,
                    'building': bld,
                    'type': l_type,
                    'url': url,
                    'isAdditional': is_additional,
                    'weekStart': w_start,
                    'weekEnd': w_end,
                    'weekType': w_type,
                    'subgroup': subgrp
                })

    # Inherit teacher and room for alternating even/odd paired lessons in cell (e.g. Demin)
    if len(results) >= 2:
        for i in range(len(results) - 1):
            r1 = results[i]
            r2 = results[i+1]
            if not r1['teacher'] and r2['teacher']:
                if (r1['weekType'] in ['even', 'odd'] and r2['weekType'] in ['even', 'odd']) or (r1['weekStart'] == 1 and r2['weekStart'] == 2):
                    r1['teacher'] = r2['teacher']
                    if not r1['room']:
                        r1['room'] = r2['room']
                        r1['building'] = r2['building']

    return results

def clean_major_and_note(raw: str) -> tuple:
    if not raw:
        return "", ""
    practice_note = ""
    practice_match = re.search(r'\(.*?(?:производственная\s+практика|практика).*?\)', raw, re.I)
    if practice_match:
        practice_note = practice_match.group(0).strip('() ')
    lines = [l.strip() for l in raw.split("\n") if l.strip()]
    first_line = lines[0] if lines else raw
    cutoff = re.search(r'(?i)\s+(?:[cс]\s+\d+|\d{1,2}[./]\d{1,2}|\d+(?:-\d+)?\s*н\b|\(\s*\d{1,2}[./]|\(\s*\d+(?:-\d+)?\s*н|\(.*практика)', first_line)
    if cutoff:
        major = first_line[:cutoff.start()]
    else:
        major = first_line
    major = re.sub(r'\(.*?\)', '', major)
    major = re.sub(r'\s+', ' ', major).strip(' ,;.-')
    if major:
        major = major[0].upper() + major[1:]
    return major, practice_note

def parse_schedule_xlsx(file_path, source_url):
    wb = openpyxl.load_workbook(file_path, data_only=True)
    result_groups = []
    lesson_counter = 0

    print(f"Loaded workbook. Sheets: {wb.sheetnames}")

    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        s_low = sheet_name.lower()

        # Determine course label
        if "магистр" in s_low:
            if "1" in s_low or "перв" in s_low:
                course_name = "1 курс (магистратура)"
            else:
                course_name = "2 курс (магистратура)"
        elif "1" in s_low:
            course_name = "1 курс"
        elif "2" in s_low:
            course_name = "2 курс"
        elif "3" in s_low:
            course_name = "3 курс"
        elif "4" in s_low:
            course_name = "4 курс"
        elif "5" in s_low:
            course_name = "5 курс"
        elif "6" in s_low:
            course_name = "6 курс"
        else:
            course_name = sheet_name

        # 1. Expand merged cells
        merged_map = {}
        for rng in ws.merged_cells.ranges:
            top_val = ws.cell(rng.min_row, rng.min_col).value
            if top_val is not None:
                for r in range(rng.min_row, rng.max_row + 1):
                    for c in range(rng.min_col, rng.max_col + 1):
                        merged_map[(r, c)] = top_val

        def get_val(r, c):
            if (r, c) in merged_map:
                return str(merged_map[(r, c)] or "").strip()
            return str(ws.cell(r, c).value or "").strip()

        # 2. Locate header row with 'день' and 'время'
        day_col = None
        time_col = None
        header_row = None

        for r in range(1, min(12, ws.max_row + 1)):
            for c in range(1, min(10, ws.max_column + 1)):
                val = get_val(r, c).lower()
                if "день" in val and day_col is None:
                    day_col = c
                    header_row = r
                elif "время" in val and time_col is None:
                    time_col = c
                    if header_row is None:
                        header_row = r

        if not header_row or not day_col or not time_col:
            print(f"Skipping sheet {sheet_name}: header not found (h={header_row}, d={day_col}, t={time_col})")
            continue

        # 3. Locate groups in header_row
        from collections import Counter
        groups = []
        for c in range(max(day_col, time_col) + 1, ws.max_column + 1):
            g_val = get_val(header_row, c).strip()
            m = re.search(r"06-\s*(\d{3})", g_val)
            if m:
                clean_g = f"06-{m.group(1)}"
                raw_major = get_val(header_row + 1, c).strip()
                clean_m, p_note = clean_major_and_note(raw_major)

                # Check if group has subgroup in parens
                sub_match = re.search(r"\((\d+)\)", g_val)
                subgrp = sub_match.group(1) if sub_match else ""
                full_group_name = f"{clean_g} ({subgrp})" if subgrp else clean_g

                groups.append({
                    "col": c,
                    "group": full_group_name,
                    "shortGroup": clean_g,
                    "major": clean_m,
                    "course": course_name,
                    "specialization": "",
                    "practiceNote": p_note
                })

        # Disambiguate duplicate group numbers on the same sheet (e.g. Master's profiles)
        counts = Counter(g["group"] for g in groups)
        for g in groups:
            if counts[g["group"]] > 1:
                g["specialization"] = g["major"]
                g["group"] = f"{g['shortGroup']} ({g['major']})"
            g["id"] = g["group"]

        print(f"Sheet '{sheet_name}': detected {len(groups)} groups for course '{course_name}'")

        # 4. Parse grid
        cur_day_name = "понедельник"
        cur_day_idx = 1
        cur_time = "08:30-10:00"
        cur_start = "08:30"
        cur_end = "10:00"

        # Initialize days map keyed by column index to isolate groups with shared numbers
        group_days_map = {
            g["col"]: {di: {"name": dn, "idx": di, "lessons": []} for dn, di in DAYS_ORDER}
            for g in groups
        }
        seen_keys = set()

        for r in range(header_row + 2, ws.max_row + 1):
            d_val = get_val(r, day_col).lower()
            t_val = get_val(r, time_col)

            for dn, di in DAYS_ORDER:
                if dn in d_val:
                    cur_day_name = dn
                    cur_day_idx = di
                    break

            time_matches = re.findall(r"(\d{1,2})[.:](\d{2})", t_val)
            if len(time_matches) >= 2:
                h1, m1 = int(time_matches[0][0]), time_matches[0][1]
                h2, m2 = int(time_matches[1][0]), time_matches[1][1]
                cur_start = f"{h1:02d}:{m1}"
                cur_end = f"{h2:02d}:{m2}"
                cur_time = f"{cur_start}-{cur_end}"

            if not cur_day_name or not cur_time:
                continue

            for g in groups:
                cell_text = get_val(r, g["col"])
                if not cell_text or cell_text.lower() in ["день", "время", "понедельник", "вторник", "среда", "четверг", "пятница", "суббота"]:
                    continue

                parsed_list = parse_physics_cell(cell_text)
                if not parsed_list:
                    continue

                for p in parsed_list:
                    key = f"{g['col']}_{cur_day_idx}_{cur_time}_{p['subject']}_{p['weekStart']}_{p['weekEnd']}_{p['weekType']}_{p['teacher']}_{p['room']}"
                    if key in seen_keys:
                        continue
                    seen_keys.add(key)

                    lesson_counter += 1
                    lesson = {
                        "id": f"{g['id']}-{lesson_counter}",
                        "subject": p["subject"],
                        "rawText": p.get("rawText", ""),
                        "teacher": p["teacher"],
                        "room": p["room"],
                        "building": p["building"],
                        "type": p["type"],
                        "url": p["url"],
                        "isAdditional": p["isAdditional"],
                        "weekType": p["weekType"],
                        "weekStart": p["weekStart"],
                        "weekEnd": p["weekEnd"],
                        "time": cur_time,
                        "timeStart": cur_start,
                        "timeEnd": cur_end
                    }
                    group_days_map[g["col"]][cur_day_idx]["lessons"].append(lesson)

        # Assemble groups for this sheet
        for g in groups:
            days_list = []
            for dn, di in DAYS_ORDER:
                d = group_days_map[g["col"]][di]
                if d["lessons"]:
                    days_list.append({
                        "dayName": d["name"],
                        "dayIndex": d["idx"],
                        "lessons": d["lessons"]
                    })

            result_groups.append({
                "id": g["id"],
                "group": g["group"],
                "course": g["course"],
                "major": g["major"],
                "days": days_list,
                "shortGroup": g["shortGroup"],
                "specialization": g["specialization"],
                "practiceNote": g.get("practiceNote") or ""
            })

    iso_date = datetime.now(timezone.utc).strftime("%Y-%m-%dT%H:%M:%SZ")
    return {
        "version": 4,
        "semester": compute_semester_string(),
        "updatedAt": iso_date,
        "sourceUrl": source_url,
        "groups": result_groups
    }

def main():
    latest_url = get_latest_xlsx_url()
    if not latest_url:
        print("Failed to find latest XLSX URL for Physics")
        sys.exit(1)

    print(f"Latest Physics XLSX URL: {latest_url}")
    local_file = os.path.join(os.path.dirname(os.path.abspath(__file__)), "physics_schedule_temp.xlsx")
    print("Downloading XLSX...")
    if not download_file(latest_url, local_file, HEADERS):
        print("Failed to download XLSX", file=sys.stderr)
        sys.exit(1)

    try:
        db = parse_schedule_xlsx(local_file, latest_url)
        save_database_with_cache_check(db, OUTPUT_FILE)
    finally:
        if os.path.exists(local_file):
            os.remove(local_file)

if __name__ == "__main__":
    main()
