import sys
import os
import re
import json
from datetime import datetime, timezone
import openpyxl

from base_parser import (
    HEADERS,
    DAYS_ORDER,
    get_latest_xlsx_url as base_get_latest_xlsx_url,
    download_file,
    compute_semester_string,
    save_database_with_cache_check
)

KFU_URL = "https://kpfu.ru/computing-technology/raspisanie"
OUTPUT_FILE = "schedule.json"

def get_latest_xlsx_url():
    return base_get_latest_xlsx_url(KFU_URL, HEADERS)

def split_cell_into_lesson_texts(text):
    trimmed = text.strip()
    if not trimmed:
        return []
    split_pattern = r"(?:;\s*(?=\()|\n\s*(?=\()|(?<=[\.\)])\s+(?=\((?:[нч]/н|(?:с\s+)?[0-9]{1,2}[^)]*?нед|[0-9]{1,2}\s*-\s*[0-9]{1,2})))"
    parts = re.split(split_pattern, trimmed, flags=re.IGNORECASE)
    res = [p.strip() for p in parts if p.strip()]
    return res if res else [trimmed]

def parse_lesson_text(txt):
    trimmed = txt.strip()
    if not trimmed:
        return None
    
    lower = trimmed.lower()
    # Reject footers and signatures
    if "____" in trimmed or \
       "департамент образования" in lower or \
       "учебно-методическ" in lower or \
       "утверждаю" in lower or \
       "согласовано" in lower or \
       "яббаров" in lower or \
       ("ивмиит" in lower and "хафизова" in lower):
        return None

    week_type = "all"
    week_start = 1
    week_end = 18

    range_pattern = r"\(\s*(?:с\s+)?(?:([нч]/н)\s*)?([0-9]{1,2})\s*-+\s*([0-9]{1,2})\s*(?:нед|недел|лаб|лек|пр|\))"
    parity_match = re.search(range_pattern, trimmed, flags=re.IGNORECASE)
    if parity_match:
        tag = parity_match.group(1)
        if tag:
            tag_low = tag.lower()
            if "н/н" in tag_low:
                week_type = "odd"
            elif "ч/н" in tag_low:
                week_type = "even"
        
        start_str = parity_match.group(2)
        end_str = parity_match.group(3)
        if start_str and end_str:
            week_start = int(start_str)
            week_end = int(end_str)
            if week_type == "all":
                if week_start == 1 and week_end == 9:
                    week_type = "first_half"
                elif week_start == 10 and week_end == 18:
                    week_type = "second_half"
    else:
        # Match single-week, comma-separated weeks or multiple ranges inside parentheses with "нед" / "недел"
        # e.g. (18 нед.), (9 нед. ), (11 нед. ), (5,7 нед. ), (16,17 нед. Пр.), (1-3,5-18 нед.)
        week_paren_match = re.search(r"\(\s*(?:с\s+)?([^)]*?(?:нед|недел)[^)]*?)\)", trimmed, flags=re.IGNORECASE)
        if week_paren_match:
            content = week_paren_match.group(1)
            content_low = content.lower()
            if "н/н" in content_low:
                week_type = "odd"
            elif "ч/н" in content_low:
                week_type = "even"
            
            nums = [int(n) for n in re.findall(r"(?<!\d)([1-9]|1[0-9]|2[0-5])(?!\d)", content)]
            if nums:
                week_start = min(nums)
                week_end = max(nums)
                if week_type == "all":
                    if week_start == 1 and week_end == 9:
                        week_type = "first_half"
                    elif week_start == 10 and week_end == 18:
                        week_type = "second_half"
        elif re.search(r"\bн/н\b|\(н/н\)", trimmed, flags=re.IGNORECASE):
            week_type = "odd"
        elif re.search(r"\bч/н\b|\(ч/н\)", trimmed, flags=re.IGNORECASE):
            week_type = "even"

    # Also check if leading "н/н" or "ч/н" was outside parenthesis, e.g. "ч/н (2-18 неделя)"
    if week_type == "all":
        if re.search(r"^\s*н/н\b", trimmed, flags=re.IGNORECASE):
            week_type = "odd"
        elif re.search(r"^\s*ч/н\b", trimmed, flags=re.IGNORECASE):
            week_type = "even"

    clean = re.sub(r"^\s*(?:(?:с|до)\s+\d{1,2}[.:]\d{2}\s*|\d{1,2}[.:]\d{2}\s*-\s*\d{1,2}[.:]\d{2}\s*|[нч]/н\s+)*\(\s*(?:с\s+)?[^\)]*(?:нед|недел|[нч]/н|\d+\s*-\s*\d+)[^\)]*\)\s*(?:нед\.?)?\s*", "", trimmed, flags=re.IGNORECASE).strip()

    room = ""
    building = ""

    sub_cleans = [s.strip() for s in clean.split(";") if s.strip()]
    subgroups = []
    
    teacher_pattern = r"[А-ЯЁ][а-яё]+\s+[А-ЯЁ]\.\s*[А-ЯЁ]\.?"

    for sub in sub_cleans:
        building = ""
        room = ""

        # 1. Building
        kreml_match = re.search(r"(?:\(\s*)?[Кк]р(?:ем|ме)[л\.]*\s*(\d+[\s\-]*(?:[А-Яа-яA-Za-z])?)\)?", sub, flags=re.IGNORECASE)
        spartak_match = re.search(r"(?:\(\s*)?[Сс]партаковская\s*(\d+[\s\-]*(?:[А-Яа-яA-Za-z])?)\)?", sub, flags=re.IGNORECASE)

        if kreml_match:
            num = kreml_match.group(1).replace(" ", "").replace("-", "").upper()
            building = f"Кремл. {num}"
        elif spartak_match:
            num = spartak_match.group(1).replace(" ", "").replace("-", "").upper()
            building = f"Спартаковская {num}"
        elif "уникс" in sub.lower() or "нужина" in sub.lower():
            building = "УНИКС"

        # 2. Room
        room_pattern = r"(?:ауд\.?\s*)+(.*?)(?=\s*(?:\(?\s*[Кк]р(?:ем|ме)[л\.]*|\(?\s*[Сс]партаковская|\(?\s*УНИКС|\(?\s*[Нн]ужина|\s*\-\s*д\.|\s*\-\s*лек\.|\s*\-\s*лаб\.|\s*\-\s*пр\.|\;|\)|$))"
        r_match = re.search(room_pattern, sub, flags=re.IGNORECASE)
        if r_match:
            r_str = r_match.group(1).strip(" ,.-;()\\t\\n")
            if r_str:
                room = r_str

        if building == "УНИКС" or "уникс" in sub.lower() or "спорткомплекс" in sub.lower():
            room = "спорткомплекс"
        elif not room and building:
            f_match = re.search(r"\b(\d{3,4}[А-Яа-яA-Za-z]?)\b", sub)
            if f_match:
                candidate = f_match.group(1).strip()
                if candidate not in ["2024", "2025", "2026", "2027"]:
                    room = candidate

        # 3. Teachers
        teachers = re.findall(teacher_pattern, sub)
        room_list = [r.strip() for r in re.split(r'[,/]', room) if r.strip()] if room else []

        if len(teachers) > 1 and len(room_list) > 1 and len(teachers) == len(room_list):
            for t, r in zip(teachers, room_list):
                subgroups.append({
                    "teacher": t,
                    "room": r,
                    "building": building
                })
        else:
            teacher_str = ", ".join(teachers)
            subgroups.append({
                "teacher": teacher_str,
                "room": room,
                "building": building
            })

    # If there are multiple subgroups, filter out completely empty ones
    if len(subgroups) > 1:
        subgroups = [sg for sg in subgroups if sg["teacher"] or sg["room"]]
    if not subgroups:
        subgroups = [{"teacher": "", "room": "", "building": ""}]

    # 4. Lesson Type
    clean_low = clean.lower()
    is_additional = bool(re.search(r"\s*-\s*д\.?(?:\s|$)", clean_low))

    if "лаб." in clean_low or "лаборатор" in clean_low:
        l_type = "lab"
    elif "дистанцион" in clean_low:
        l_type = "distance"
    elif "эор" in clean_low:
        l_type = "eor"
    elif "цор" in clean_low:
        l_type = "cor"
    elif "практи" in clean_low or "пр." in clean_low:
        l_type = "practice"
    elif "лек." in clean_low or "лекция" in clean_low:
        l_type = "lecture"
    else:
        l_type = "other"

    # 5. Subject
    subject = clean
    room_rx = re.search(r"(?:ауд\.?|\b[0-9]{3,4}\s*\(?\s*[Кк]р(?:ем|ме)[л\.]*)", clean, flags=re.IGNORECASE)
    if room_rx:
        subject = subject[:room_rx.start()]
    
    for bp in ["шахматный центр", "кск кфу"]:
        idx = subject.lower().find(bp)
        if idx != -1:
            subject = subject[:idx]
    
    all_teachers = re.findall(teacher_pattern, clean)
    for t in all_teachers:
        subject = subject.replace(t, "")
    
    subject = re.sub(r"\s*-\s*д\.?(?:\s|$)", "", subject, flags=re.IGNORECASE)
    subject = subject.split(";")[0]

    # Clean format/EOR/distance markers
    subject = re.sub(r"\s*-\s*(?:лекция|занятие\s+проводится|в)\s*(?:в\s+)?(?:дистанционном\s+формате|формате\s+[ЭэЦц][Оо][Рр]|формате\s+ДО|[ЭэЦц][Оо][Рр]).*", "", subject, flags=re.IGNORECASE)
    subject = re.sub(r"\s*-\s*в\s+формате.*", "", subject, flags=re.IGNORECASE)
    subject = re.sub(r"\s*\(?\bв\s+формате\s+[ЭэЦц][Оо][Рр]\)?", "", subject, flags=re.IGNORECASE)
    subject = re.sub(r"\s*-\s*(?:дистанционн|в\s+дистанционн).*", "", subject, flags=re.IGNORECASE)

    subject = subject.strip(" ,.-;:\t\n")
    if subject.endswith(")") and subject.count(")") > subject.count("("):
        subject = subject[:-1].strip(" ,.-;:\t\n")
    if subject.startswith("(") and subject.count("(") > subject.count(")"):
        subject = subject[1:].strip(" ,.-;:\t\n")
    if subject.startswith(")") and subject.count(")") > subject.count("("):
        subject = subject[1:].strip(" ,.-;:\t\n")
    subject = re.sub(r"\s+", " ", subject)

    # 6. Custom time (PE)
    custom_time = None
    time_rx = re.search(r"(\d{1,2})[.:](\d{2})\s*-\s*(\d{1,2})[.:](\d{2})", trimmed)
    if time_rx:
        low_trimmed = trimmed.lower()
        if "проводится" in low_trimmed or "спорт" in low_trimmed or "физическ" in low_trimmed or "уникс" in low_trimmed:
            h1 = int(time_rx.group(1))
            m1 = time_rx.group(2)
            h2 = int(time_rx.group(3))
            m2 = time_rx.group(4)
            if 7 <= h1 <= 21 and 7 <= h2 <= 22:
                custom_time = {
                    "start": f"{h1:02d}:{m1}",
                    "end": f"{h2:02d}:{m2}"
                }

    # 7. URL extraction
    lesson_url = ""
    url_match = re.search(r"https?://[^\s,\);]+", trimmed)
    if url_match:
        lesson_url = url_match.group(0).strip(" .,;")
    elif "edu.kpfu.ru" in trimmed.lower():
        lesson_url = "https://edu.kpfu.ru"

    results = []
    for sg in subgroups:
        t_type = l_type
        if t_type == "other":
            r_str = sg["room"]
            b_str = sg["building"]
            m_num = re.search(r"\d+", r_str)
            if m_num:
                r_num = int(m_num.group(0))
                if "35" in b_str:
                    # Floors <= 2: 1xx, 2xx (108, 109, 216, 218) are lecture halls
                    if r_num < 300:
                        t_type = "lecture"
                    else:
                        t_type = "practice"
                elif "16" in b_str:
                    # Floor 1: 1xx (110, 112) are lecture halls; 200+ are practice/labs
                    if 100 <= r_num < 200:
                        t_type = "lecture"
                    else:
                        t_type = "practice"
                elif "спартак" in b_str.lower():
                    t_type = "practice"
            elif b_str == "УНИКС" or "спорт" in clean_low or "культур" in clean_low:
                t_type = "practice"

        results.append({
            "subject": subject,
            "clean": clean,
            "teacher": sg["teacher"],
            "room": sg["room"],
            "building": sg["building"],
            "type": t_type,
            "url": lesson_url,
            "isAdditional": is_additional,
            "weekType": week_type,
            "weekStart": week_start,
            "weekEnd": week_end,
            "customTime": custom_time
        })

    return results

def parse_schedule_xlsx(file_path, source_url):
    print("Loading workbook...")
    if openpyxl is not None:
        wb = openpyxl.load_workbook(file_path, data_only=True)
        sheet = wb.active

        # Fast cell getter resolving merged ranges
        merged_dict = {}
        for mr in sheet.merged_cells.ranges:
            val = sheet.cell(row=mr.min_row, column=mr.min_col).value
            if val is not None:
                s_val = str(val).strip()
                if s_val:
                    for r in range(mr.min_row, mr.max_row + 1):
                        for c in range(mr.min_col, mr.max_col + 1):
                            merged_dict[(r, c)] = s_val

        def get_cell_val(row, col):
            if (row, col) in merged_dict:
                return merged_dict[(row, col)]
            v = sheet.cell(row=row, column=col).value
            return str(v).strip() if v is not None else ""

        max_col = sheet.max_column
        max_row = sheet.max_row
    else:
        import zipfile
        import xml.etree.ElementTree as ET

        cells_dict = {}
        merged_dict = {}
        with zipfile.ZipFile(file_path) as z:
            shared_strings = []
            if "xl/sharedStrings.xml" in z.namelist():
                tree = ET.fromstring(z.read("xl/sharedStrings.xml"))
                ns = {"ns": "http://schemas.openxmlformats.org/spreadsheetml/2006/main"}
                for si in tree.findall("ns:si", ns):
                    text = "".join(t.text for t in si.findall(".//ns:t", ns) if t.text)
                    shared_strings.append(text)

            sheet_tree = ET.fromstring(z.read("xl/worksheets/sheet1.xml"))
            ns = {"ns": "http://schemas.openxmlformats.org/spreadsheetml/2006/main"}

            def cell_ref_to_row_col(ref):
                m = re.match(r"([A-Z]+)(\d+)", ref)
                col_str, row_str = m.group(1), m.group(2)
                col = 0
                for char in col_str:
                    col = col * 26 + (ord(char) - ord("A") + 1)
                return int(row_str), col

            for c_elem in sheet_tree.findall(".//ns:c", ns):
                r, c = cell_ref_to_row_col(c_elem.attrib["r"])
                t = c_elem.attrib.get("t")
                v_elem = c_elem.find("ns:v", ns)
                if v_elem is not None and v_elem.text is not None:
                    val = v_elem.text
                    if t == "s":
                        val = shared_strings[int(val)]
                    cells_dict[(r, c)] = str(val).strip()

            for mc_elem in sheet_tree.findall(".//ns:mergeCell", ns):
                ref = mc_elem.attrib["ref"]
                start_ref, end_ref = ref.split(":")
                r1, c1 = cell_ref_to_row_col(start_ref)
                r2, c2 = cell_ref_to_row_col(end_ref)
                top_val = cells_dict.get((r1, c1), "")
                if top_val:
                    for r in range(r1, r2 + 1):
                        for c in range(c1, c2 + 1):
                            merged_dict[(r, c)] = top_val

        def get_cell_val(row, col):
            if (row, col) in merged_dict:
                return merged_dict[(row, col)]
            return cells_dict.get((row, col), "")

        max_col = max(c for (r, c) in cells_dict.keys())
        max_row = max(r for (r, c) in cells_dict.keys())
    start_col = 5  # Column 'E'

    groups = []
    cur_course = "1 курс"
    cur_major = ""

    print("Scanning groups...")
    for ci in range(start_col, max_col + 1):
        r17 = get_cell_val(17, ci)
        r17_low = r17.lower()
        if "магистратура" in r17_low:
            if "1" in r17_low or "первый" in r17_low:
                cur_course = "1 курс (магистратура)"
            elif "2" in r17_low or "второй" in r17_low:
                cur_course = "2 курс (магистратура)"
            else:
                cur_course = "Магистратура"
        elif "курс" in r17_low:
            cur_course = r17.strip()
        elif r17:
            cur_major = re.sub(r'\s+', ' ', r17).strip()

        g_name = get_cell_val(18, ci)
        if g_name and "курс" not in g_name.lower() and re.search(r'(?:\d{2}-)?\d{3}', g_name):
            final_name = g_name.strip()
            if not final_name.startswith("09-"):
                final_name = f"09-{final_name}"

            course_name = cur_course
            is_master = ".04." in cur_major or "магистр" in cur_major.lower() or "магистр" in cur_course.lower()
            if is_master:
                if "09-6" in final_name:
                    course_name = "1 курс (магистратура)"
                elif "09-5" in final_name:
                    course_name = "2 курс (магистратура)"
                else:
                    course_name = "Магистратура"

            m_base = re.match(r"^(09-\d{3})", final_name)
            base = m_base.group(1) if m_base else final_name
            m_sub = re.search(r"\((\d+)\)", final_name)
            subgrp = m_sub.group(1) if m_sub else ""

            rem = final_name.replace(base, "")
            if subgrp:
                rem = re.sub(r"\(" + subgrp + r"\)", "", rem)
            spec = rem.strip(" ()")

            if is_master:
                c_num = "1 курс" if "09-6" in base else "2 курс"
                if subgrp:
                    clean_name = f"{base} ({subgrp}) ({c_num})"
                    short_name = f"{base} ({subgrp})"
                else:
                    clean_name = f"{base} ({c_num})"
                    short_name = base
            else:
                if subgrp:
                    clean_name = f"{base} ({subgrp})"
                    short_name = f"{base} ({subgrp})"
                else:
                    clean_name = base
                    short_name = base

            groups.append({
                "colIdx": ci,
                "course": course_name,
                "major": cur_major,
                "group": clean_name,
                "shortGroup": short_name,
                "specialization": spec
            })

    print(f"Detected {len(groups)} groups.")

    days_order = [
        ("понедельник", 1), ("вторник", 2), ("среда", 3),
        ("четверг", 4), ("пятница", 5), ("суббота", 6)
    ]

    row_meta = {}
    cur_day_name = "понедельник"
    cur_day_idx = 1
    cur_time = "08:30-10:00"
    cur_start = "08:30"
    cur_end = "10:00"

    for ri in range(19, min(max_row + 1, 116)):
        c_val = get_cell_val(ri, 3).lower() # Column C
        d_val = get_cell_val(ri, 4)         # Column D

        for dn, di in days_order:
            if dn in c_val:
                cur_day_name = dn
                cur_day_idx = di
                break

        time_matches = re.findall(r"(\d{1,2})[.:](\d{2})", d_val)
        if len(time_matches) >= 2:
            h1 = int(time_matches[0][0])
            m1 = time_matches[0][1]
            h2 = int(time_matches[1][0])
            m2 = time_matches[1][1]
            cur_start = f"{h1:02d}:{m1}"
            cur_end = f"{h2:02d}:{m2}"
            cur_time = f"{cur_start}-{cur_end}"

        row_meta[ri] = {
            "dayName": cur_day_name,
            "dayIdx": cur_day_idx,
            "timeSlot": cur_time,
            "start": cur_start,
            "end": cur_end
        }

    result_groups = []
    lesson_counter = 0

    print("Parsing schedule grid...")
    for g in groups:
        ci = g["colIdx"]
        days_map = {di: {"name": dn, "idx": di, "lessons": []} for dn, di in days_order}
        seen = set()

        for ri in range(19, min(max_row + 1, 116)):
            val = get_cell_val(ri, ci)
            if not val:
                continue
            if val.lower() in ["вт", "ср", "чт", "пт", "сб", "вск"]:
                continue

            lesson_chunks = split_cell_into_lesson_texts(val)
            for chunk in lesson_chunks:
                parsed_list = parse_lesson_text(chunk)
                if not parsed_list:
                    continue
                for parsed in parsed_list:
                    meta = row_meta.get(ri)
                    if not meta:
                        continue

                    key = f"{meta['dayIdx']}_{meta['timeSlot']}_{parsed['clean']}_{parsed['weekStart']}_{parsed['weekEnd']}_{parsed['weekType']}_{parsed['teacher']}_{parsed['room']}"
                    if key in seen:
                        continue
                    seen.add(key)

                    lesson_counter += 1
                    start_slot = parsed["customTime"]["start"] if parsed["customTime"] else meta["start"]
                    end_slot = parsed["customTime"]["end"] if parsed["customTime"] else meta["end"]
                    time_slot = f"{start_slot}-{end_slot}" if parsed["customTime"] else meta["timeSlot"]

                    lesson = {
                        "id": f"{g['group']}-{lesson_counter}",
                        "subject": parsed["subject"],
                        "rawText": parsed["clean"],
                        "teacher": parsed["teacher"],
                        "room": parsed["room"],
                        "building": parsed["building"],
                        "type": parsed["type"],
                        "url": parsed["url"],
                        "isAdditional": parsed["isAdditional"],
                        "weekType": parsed["weekType"],
                        "weekStart": parsed["weekStart"],
                        "weekEnd": parsed["weekEnd"],
                        "time": time_slot,
                        "timeStart": start_slot,
                        "timeEnd": end_slot
                    }
                    days_map[meta["dayIdx"]]["lessons"].append(lesson)

        days_list = []
        for dn, di in days_order:
            d = days_map[di]
            if d["lessons"]:
                days_list.append({
                    "dayName": d["name"],
                    "dayIndex": d["idx"],
                    "lessons": d["lessons"]
                })

        result_groups.append({
            "id": g["group"],
            "group": g["group"],
            "course": g["course"],
            "major": g["major"],
            "days": days_list,
            "shortGroup": g["shortGroup"],
            "specialization": g["specialization"]
        })

    iso_date = datetime.now(timezone.utc).strftime("%Y-%m-%dT%H:%M:%SZ")
    return {
        "version": 4,
        "semester": compute_semester_string(),
        "updatedAt": iso_date,
        "sourceUrl": source_url,
        "groups": result_groups
    }

def main():
    latest_url = get_latest_xlsx_url()
    if not latest_url:
        print("Failed to find latest XLSX URL")
        sys.exit(1)
    
    print(f"Latest XLSX URL: {latest_url}")
    local_file = "schedule.xlsx"
    print("Downloading XLSX...")
    if not download_file(latest_url, local_file, HEADERS):
        print("Failed to download XLSX", file=sys.stderr)
        sys.exit(1)
    
    try:
        db = parse_schedule_xlsx(local_file, latest_url)
        save_database_with_cache_check(db, OUTPUT_FILE)
    finally:
        if os.path.exists(local_file):
            os.remove(local_file)

if __name__ == "__main__":
    main()
