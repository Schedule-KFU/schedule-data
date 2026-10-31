import sys
import os
import re
import json
import tempfile
from collections import defaultdict, deque
from datetime import datetime, timezone
import openpyxl

from base_parser import (
    HEADERS,
    DAYS_ORDER,
    compute_semester_string,
    compute_semester_meta
)

try:
    import requests
except ImportError:
    requests = None

IFMIB_SCHEDULE_URL = "https://kpfu.ru/biology-medicine/uchebnyj-process/raspisanie-zachjotov-i-ekzamenov"
OUTPUT_FILE = "schedule_ifmib.json"

TEACHER_PATTERN = r"[А-ЯЁ][а-яё]+(?:-[А-ЯЁ][а-яё]+)?\s+[А-ЯЁ]\.(?:\s*[А-ЯЁ]\.?)?"
WEEK_ITEM = r'(?:\b(?:[1-9]|1[0-9]|2[0-4])(?:\s*-\s*(?:[1-9]|1[0-9]|2[0-4]))?\b)'
WEEK_LIST = rf'(?:{WEEK_ITEM}(?:\s*,\s*{WEEK_ITEM})*)'
WEEK_PREFIX = r'(?:(?<=[\(\),])|(?<=\))\s*|(?<=\bлекций\b)\s*|(?<=\bлекции\b)\s*|(?<=\bлекция\b)\s*|\A\s*)'
WEEK_SEG_PATTERN = rf'{WEEK_PREFIX}\s*({WEEK_LIST}\s*нед\.?)\s*({TEACHER_PATTERN})'


def extract_ifmib_room_building(text):
    """Extract room and building identifiers from text."""
    building = ''
    room = ''

    if re.search(r'к\.?\s*маркса\s*,?\s*74а|км\s*,?\s*74а|сим\.?\s*центр', text, re.I):
        building = 'К.Маркса 74а (сим.центр)'
    elif re.search(r'к\.?\s*маркса\s*,?\s*76|км\s*,?\s*76', text, re.I):
        building = 'К.Маркса 76 корп. 2'
    elif re.search(r'к\.?\s*маркса\s*,?\s*74|км\s*,?\s*74', text, re.I):
        building = 'К.Маркса 74'
    elif re.search(r'лобачевского\s*1\s*/\s*29', text, re.I):
        building = 'Лобачевского 1/29'
    elif re.search(r'оренб[а-я\.]*\s*тракт\s*,?\s*10', text, re.I):
        building = 'Оренб. тракт 10'
    elif re.search(r'шмидта\s*,?\s*31', text, re.I):
        building = 'Шмидта 31'
    elif re.search(r'ершова\s*,?\s*2', text, re.I):
        building = 'Ершова 2'
    elif re.search(r'дркб', text, re.I):
        building = 'ДРКБ (терминал 10)'
    elif re.search(r'ркод|сибирск[а-я\.]*\s*тракт\s*,?\s*31', text, re.I):
        building = 'РКОД (Сибирский тракт 31)'
    elif re.search(r'здоровье\s*семьи|2-?я\s*юго-?западная', text, re.I):
        building = 'Здоровье семьи (2-я Юго-Западная 35)'
    elif re.search(r'гб\s*№?\s*7', text, re.I):
        building = 'ГБ № 7'
    elif re.search(r'бустан', text, re.I):
        building = 'СК Бустан'
    elif re.search(r'уникс', text, re.I):
        building = 'УНИКС'
    elif re.search(r'кремл[а-я\.]*\s*,?\s*16а', text, re.I):
        building = 'Кремл. 16А'
    elif re.search(r'кремл[а-я\.]*\s*,?\s*18', text, re.I):
        building = 'Кремл. 18'
    elif re.search(r'кремл[а-я\.]*\s*,?\s*35', text, re.I):
        building = 'Кремл. 35'

    r_m = re.search(r'\b(?:ауд\.?)\s*([0-9A-Za-zА-Яа-я\/\-]+)', text, re.I)
    if r_m:
        room = r_m.group(1).strip(' ,.-;')
    elif re.search(r'учеб\.?\s*ком\.?\s*(\d+)', text, re.I):
        m_uk = re.search(r'учеб\.?\s*ком\.?\s*(\d+)', text, re.I)
        room = f'учеб. комната {m_uk.group(1)}'
    elif 'учеб.ком' in text.lower() or 'учеб. ком' in text.lower():
        room = 'учеб. комната'
    elif 'бустан' in text.lower():
        room = 'спорткомплекс'

    return room, building


def split_multiteacher_chunk(chunk):
    """Split a single chunk if it defines multiple sequential or alternating teacher segments with week lists."""
    matches = list(re.finditer(WEEK_SEG_PATTERN, chunk))
    if len(matches) < 2:
        return [chunk]

    first_m = matches[0]
    head = chunk[:first_m.start()].strip()

    l_type_str = ''
    if re.search(r'лекци|лек\.', chunk, re.I):
        l_type_str = '(лекция)'
    elif re.search(r'лаб', chunk, re.I):
        l_type_str = '(лаб.)'

    head_clean = re.sub(r'\(\s*\d+\s*лекци[йия][\)]?', ' ', head, flags=re.I)
    head_clean = re.sub(r'\(\s*лекци[йия]\s*\)', ' ', head_clean, flags=re.I)
    head_clean = re.sub(r'\(\s*лаб\.?\s*\)', ' ', head_clean, flags=re.I)
    head_clean = re.sub(r'\b(?:лекция|лекций|практика|лаб\.?)\b', ' ', head_clean, flags=re.I)
    head_clean = head_clean.strip(' ,.-;:()')

    g_room, g_building = extract_ifmib_room_building(chunk)
    g_room_str = f"ауд. {g_room}" if g_room else ""
    g_loc_str = f"{g_room_str} {g_building}".strip()

    sub_chunks = []
    for i, m in enumerate(matches):
        next_start = matches[i+1].start() if i + 1 < len(matches) else len(chunk)
        tail = chunk[m.end():next_start]
        r, b = extract_ifmib_room_building(tail)
        if r or b:
            r_str = f"ауд. {r}" if r else ""
            loc_str = f"{r_str} {b}".strip()
        else:
            loc_str = g_loc_str

        week_str = f"({m.group(1).strip()})"
        teacher_str = m.group(2).strip()

        parts = [head_clean]
        if l_type_str:
            parts.append(l_type_str)
        parts.append(week_str)
        parts.append(teacher_str)
        if loc_str:
            parts.append(loc_str)
        sub_chunk = " ".join(parts)
        sub_chunks.append(sub_chunk)

    return sub_chunks


def get_ifmib_xlsx_urls(retries=3):
    """Scrape the IFMiB schedule page to dynamically discover active 1-3 course XLSX URLs."""
    html = ""
    for attempt in range(retries):
        try:
            if requests is not None:
                resp = requests.get(IFMIB_SCHEDULE_URL, headers=HEADERS, timeout=35)
                resp.encoding = "windows-1251"
                html = resp.text
            else:
                import urllib.request
                req = urllib.request.Request(IFMIB_SCHEDULE_URL, headers=HEADERS)
                with urllib.request.urlopen(req, timeout=35) as resp:
                    html = resp.read().decode("windows-1251", errors="replace")
            if html:
                break
        except Exception as e:
            print(f"[Warning] Attempt {attempt+1}/{retries} to fetch {IFMIB_SCHEDULE_URL} failed: {e}", file=sys.stderr)
            if attempt < retries - 1:
                import time
                time.sleep(2)

    if not html:
        return []

    matches = re.findall(r'<a[^>]+href=[\"\']([^\"\']+)[\"\'][^>]*>(.*?)</a>', html, re.DOTALL | re.IGNORECASE)
    selected = []
    seen_urls = set()

    for href, raw_text in matches:
        text = re.sub(r'<[^>]+>', '', raw_text).replace('&nbsp;', ' ').strip()
        href_low = href.lower()
        text_low = text.lower()

        if 'ias_utils.file_download' in href and href_low.endswith('.xlsx'):
            if any(ex in href_low or ex in text_low for ex in ['dop', 'komiss', 'sessii', 'сесси', 'пересдач', 'lekcionnyj']):
                continue
            if href in seen_urls:
                continue

            m_kurs = re.search(r'\b([1-3])\s*(?:kurs|курс|course)\b', text_low) or re.search(r'([1-3])\.kurs', href_low)
            if m_kurs:
                c_num = int(m_kurs.group(1))
                is_angl = 'angl' in href_low or 'course' in text_low or 'англ' in text_low
                seen_urls.add(href)
                selected.append({
                    'courseNum': c_num,
                    'course': f'{c_num} курс',
                    'isAngl': is_angl,
                    'label': f'{c_num} курс' + (' (англ)' if is_angl else ''),
                    'url': href
                })

    selected.sort(key=lambda x: (x['courseNum'], x['isAngl']))
    return selected


def split_ifmib_cell_into_lessons(text):
    """Split cell content if it contains multiple lessons (e.g. separated by slash or comma)."""
    trimmed = text.strip()
    if not trimmed:
        return []

    # Exclude signatures and footers
    low = trimmed.lower()
    if any(k in low for k in ['исполнитель', 'департамент образования', 'учебно-методическ', 'утверждаю', 'согласовано', 'яббаров', 'андреева']):
        return []

    # Split on slash before capital letter (avoids 507/508 or 1/29)
    slash_parts = re.split(r'\s*/\s*(?=[А-ЯЁ])', trimmed)
    sub_parts = []
    for sp in slash_parts:
        # Split on comma if followed by a new subject with week pattern
        comma_splits = re.split(r',\s*(?=[А-ЯЁ][а-яё\s]+(?:\([0-9\s,-]+нед))', sp)
        sub_parts.extend(comma_splits)

    results = []
    for p in sub_parts:
        p_str = p.strip()
        if not p_str:
            continue
        results.extend(split_multiteacher_chunk(p_str))

    # If first sublesson lacks room/building but a trailing sublesson has it, share it
    if len(results) > 1:
        trailing = results[-1]
        b_match = re.search(r'(?:к\.?\s*маркса|км\s*,|лобачевского|оренб|шмидта|ершова|гб|бустан|кремл)', trailing, re.I)
        if b_match:
            shared_tail = trailing[b_match.start():]
            shared_room = re.search(r'(?:ауд\.?|учеб\.ком\.?)\s*[0-9A-Za-zА-Яа-я\/\-]+', trailing, re.I)
            for idx in range(len(results) - 1):
                cur = results[idx]
                has_b = bool(re.search(r'(?:к\.?\s*маркса|км\s*,|лобачевского|оренб|шмидта|ершова|гб|бустан|кремл)', cur, re.I))
                has_r = bool(re.search(r'(?:ауд\.?|учеб\.ком\.?)\s*[0-9A-Za-zА-Яа-я\/\-]+', cur, re.I))
                if not has_b and not has_r and shared_room:
                    results[idx] = f"{cur} {shared_room.group(0)} {shared_tail}"
                elif not has_b:
                    results[idx] = f"{cur} {shared_tail}"

    return results if results else [trimmed]


def parse_ifmib_lesson_item(raw_text):
    """Parse an individual IFMiB lesson entry from raw cell text."""
    text = raw_text.strip()
    if not text:
        return None

    # 1. Custom time at start, e.g. '10.30-12.55'
    custom_time = None
    time_m = re.match(r'^(\d{1,2})[.:](\d{2})\s*-\s*(\d{1,2})[.:](\d{2})\s+', text)
    if time_m:
        h1, m1, h2, m2 = int(time_m.group(1)), time_m.group(2), int(time_m.group(3)), time_m.group(4)
        if 7 <= h1 <= 21 and 7 <= h2 <= 22:
            custom_time = {
                'start': f'{h1:02d}:{m1}',
                'end': f'{h2:02d}:{m2}'
            }
            text = text[time_m.end():].strip()

    # 2. Extract weeks and parity
    week_start = 1
    week_end = 18
    week_type = 'all'
    explicit_weeks = None

    # 2.1 First look for parenthesis containing 'нед'
    week_m = re.search(r'\((?:[^)]*?(?:нед|недел)[^)]*?)\)', text, re.IGNORECASE)
    week_info_str = ''
    if week_m:
        week_info_str = week_m.group(0)

    # 2.2 If not found in parens, search for "X, Y, Z нед." in text
    if not week_info_str:
        ext_m = re.search(rf'({WEEK_LIST}\s*нед\.?)', text, re.IGNORECASE)
        if ext_m:
            week_info_str = ext_m.group(0)

    # 2.3 If still not found, check if there's any parenthesis with numbers/range (e.g. "(1-9)")
    if not week_info_str:
        num_m = re.search(r'\((?:[^)]*?\b\d+\s*-\s*\d+\b[^)]*?)\)', text)
        if num_m and 'лекци' not in num_m.group(0).lower():
            week_info_str = num_m.group(0)

    if week_info_str:
        nums = [int(n) for n in re.findall(r'\b([1-9]|1[0-9]|2[0-5])\b', week_info_str)]
        ranges = re.findall(r'\b([1-9]|1[0-9])\s*-\s*([1-9]|1[0-9])\b', week_info_str)
        all_nums = set(nums)
        for r_s, r_e in ranges:
            all_nums.update(range(int(r_s), int(r_e) + 1))

        if all_nums:
            week_start = min(all_nums)
            week_end = max(all_nums)
            sorted_weeks = sorted(list(all_nums))
            if len(sorted_weeks) < (week_end - week_start + 1):
                explicit_weeks = sorted_weeks
                if all(w % 2 == 0 for w in sorted_weeks):
                    week_type = 'even'
                elif all(w % 2 == 1 for w in sorted_weeks):
                    week_type = 'odd'
            elif week_start == 1 and week_end == 9:
                week_type = 'first_half'
            elif week_start == 10 and week_end == 18:
                week_type = 'second_half'

    text_low = text.lower()
    if 'н/н' in text_low or 'неч' in text_low:
        week_type = 'odd'
    elif 'ч/н' in text_low or 'чет' in text_low:
        week_type = 'even'

    # 3. Lesson type
    if 'лекци' in text_low or 'лек.' in text_low:
        l_type = 'lecture'
    elif 'лаб.' in text_low or 'лаборатор' in text_low:
        l_type = 'lab'
    elif 'дистанцион' in text_low:
        l_type = 'distance'
    elif 'эор' in text_low:
        l_type = 'eor'
    elif 'цор' in text_low:
        l_type = 'cor'
    else:
        l_type = 'practice'

    # 4. Teachers
    teachers = re.findall(TEACHER_PATTERN, text)
    teacher_str = ', '.join(teachers)

    # 5. Building & Room
    room, building = extract_ifmib_room_building(text)

    # 6. Extract clean subject
    subject = text
    if week_info_str:
        subject = subject.replace(week_info_str, ' ')
    for t in teachers:
        subject = subject.replace(t, ' ')

    subject = re.sub(r'\b(?:ауд\.?)\s*[0-9A-Za-zА-Яа-я\/\-]+', ' ', subject, flags=re.I)
    subject = re.sub(r'учеб\.?\s*ком\.?\s*\d*', ' ', subject, flags=re.I)

    for b_pat in [
        r'к\.?\s*маркса\s*,?\s*74а(?:\s*\(сим\.центр\))?',
        r'к\.?\s*маркса\s*,?\s*76(?:\s*\(кор\.?2\))?',
        r'к\.?\s*маркса\s*,?\s*74',
        r'км\s*,?\s*74а(?:\s*\(сим\.центр\))?',
        r'км\s*,?\s*76(?:\s*\(кор\.?2\))?',
        r'км\s*,?\s*74',
        r'лобачевского\s*1\s*/\s*29',
        r'оренб[а-я\.]*\s*тракт\s*,?\s*10',
        r'шмидта\s*,?\s*31',
        r'ершова\s*,?\s*2',
        r'дркб(?:\s*\(10\s*тери?минал\))?',
        r'ркод',
        r'сибирск[а-я\.]*\s*тракт\s*,?\s*31',
        r'здоровье\s*семьи\s*,?',
        r'2-?я\s*юго-?западная\s*,?\s*35',
        r'гб\s*№?\s*7',
        r'ск\s*\"бустан\"',
        r'бустан',
        r'кремлевская\s*,?\s*18',
        r'кремлевская\s*,?\s*16а',
        r'кремл[а-я\.]*\s*,?\s*35',
        r'кремл[а-я\.]*\s*,?\s*16а',
        r'кремл[а-я\.]*\s*,?\s*18'
    ]:
        subject = re.sub(b_pat, ' ', subject, flags=re.I)

    subject = re.sub(r'\(кор\.?2\)|\(сим\.центр\)', ' ', subject, flags=re.I)
    subject = re.sub(r'\(\s*\d+\s*лекци[йия][\)]?', ' ', subject, flags=re.I)
    subject = re.sub(r'\b(?:лекция|лекций|практика|лаб\.?)\b', ' ', subject, flags=re.I)
    subject = re.sub(r'\(?\s*[0-9\s,-]+нед\.?\s*\)?', ' ', subject, flags=re.I)
    subject = re.sub(r'\s+', ' ', subject).strip(' ,.-;:()')

    res = {
        'subject': subject,
        'rawText': raw_text.strip(),
        'teacher': teacher_str,
        'room': room,
        'building': building,
        'type': l_type,
        'weekType': week_type,
        'weekStart': week_start,
        'weekEnd': week_end,
    }
    if explicit_weeks:
        res['weeks'] = explicit_weeks
    if custom_time:
        res['customTime'] = custom_time
    return res


def parse_schedule_sheet(wb, course_name, is_angl):
    """Parse one IFMiB workbook into groups and days."""
    sheet = wb.active

    # Fast merged cells map
    merged_map = {}
    for rng in sheet.merged_cells.ranges:
        v = sheet.cell(row=rng.min_row, column=rng.min_col).value
        if v is not None:
            for ro in range(rng.min_row, rng.max_row + 1):
                for co in range(rng.min_col, rng.max_col + 1):
                    merged_map[(ro, co)] = v

    # Dynamically locate headers
    r_spec = 12
    r_group = 13
    r_subgroup = 14
    for ro in range(1, 20):
        c2 = str(sheet.cell(row=ro, column=2).value or '').lower()
        if 'специальност' in c2:
            r_spec = ro
        if 'номер группы' in c2 or ('групп' in c2 and 'п/' not in c2):
            r_group = ro
        if 'п/групп' in c2 or 'подгрупп' in c2:
            r_subgroup = ro

    # Dynamically locate day rows and time slots (must be after r_subgroup)
    row_meta = {}
    cur_dname = "понедельник"
    cur_didx = 1
    for ro in range(r_subgroup + 1, sheet.max_row + 1):
        c1 = str(sheet.cell(row=ro, column=1).value or '').strip().lower()
        c2 = str(sheet.cell(row=ro, column=2).value or '').strip()

        for dn, di in DAYS_ORDER:
            if dn in c1:
                cur_dname = dn
                cur_didx = di
                break

        tm = re.findall(r'(\d{1,2})[.:](\d{2})', c2)
        if len(tm) >= 2:
            h1, m1 = int(tm[0][0]), tm[0][1]
            h2, m2 = int(tm[1][0]), tm[1][1]
            if 7 <= h1 <= 22 and 7 <= h2 <= 22:
                s = f'{h1:02d}:{m1}'
                e = f'{h2:02d}:{m2}'
                row_meta[ro] = {
                    'dayName': cur_dname,
                    'dayIdx': cur_didx,
                    'start': s,
                    'end': e,
                    'timeSlot': f'{s}-{e}'
                }

    # Discover groups from columns
    detected_groups = []
    for co in range(3, sheet.max_column + 1):
        spec_raw = str(merged_map.get((r_spec, co), sheet.cell(row=r_spec, column=co).value) or '').strip()
        grp_raw = str(merged_map.get((r_group, co), sheet.cell(row=r_group, column=co).value) or '').strip()
        sub_raw = str(merged_map.get((r_subgroup, co), sheet.cell(row=r_subgroup, column=co).value) or '').strip()

        if not grp_raw and not sub_raw:
            continue

        spec_note = "Англоязычная программа" if is_angl else ""
        m_note = re.search(r'\((.*?)\)', grp_raw)
        if m_note:
            note_content = m_note.group(1).strip()
            spec_note = f"{note_content}, {spec_note}" if spec_note else note_content
            grp_base = re.sub(r'\(.*?\)', '', grp_raw).strip()
        else:
            grp_base = grp_raw

        grp_base = re.sub(r'^01-2-', '01.2-', grp_base)
        if not grp_base:
            m_sub_base = re.match(r'^(01[.\-]\d+[\-]\d+)', sub_raw)
            grp_base = m_sub_base.group(1) if m_sub_base else sub_raw

        # Subgroup letter
        m_let = re.search(r'([а-яёa-z0-9])$', sub_raw.lower())
        sub_let = m_let.group(1) if m_let else ''
        if sub_let in ['а', 'a']:
            sub_let = 'а'
        elif sub_let in ['б', 'b']:
            sub_let = 'б'
        elif sub_let in ['в', 'v']:
            sub_let = 'в'

        final_group = f"{grp_base} ({sub_let})" if sub_let else grp_base

        detected_groups.append({
            'col': co,
            'group': final_group,
            'shortGroup': final_group,
            'course': course_name,
            'major': spec_raw,
            'specialization': spec_note
        })

    # Parse lessons for each detected group
    result_groups = []
    for g_info in detected_groups:
        co = g_info['col']
        days_map = {di: {'name': dn, 'idx': di, 'lessons': []} for dn, di in DAYS_ORDER}
        seen = set()

        for ro, meta in row_meta.items():
            val = merged_map.get((ro, co), sheet.cell(row=ro, column=co).value)
            if not val or not str(val).strip():
                continue

            lesson_texts = split_ifmib_cell_into_lessons(str(val))
            for chunk in lesson_texts:
                parsed = parse_ifmib_lesson_item(chunk)
                if not parsed or not parsed['subject']:
                    continue

                start_slot = parsed['customTime']['start'] if 'customTime' in parsed else meta['start']
                end_slot = parsed['customTime']['end'] if 'customTime' in parsed else meta['end']
                time_slot = f"{start_slot}-{end_slot}"

                dedup_key = (
                    meta['dayIdx'],
                    start_slot,
                    end_slot,
                    parsed['subject'].lower(),
                    parsed['teacher'].lower(),
                    parsed['weekStart'],
                    parsed['weekEnd'],
                    parsed['weekType'],
                    tuple(parsed.get('weeks') or [])
                )
                if dedup_key in seen:
                    continue
                seen.add(dedup_key)

                lesson_obj = {
                    'subject': parsed['subject'],
                    'rawText': parsed['rawText'],
                    'teacher': parsed['teacher'],
                    'room': parsed['room'],
                    'building': parsed['building'],
                    'type': parsed['type'],
                    'url': '',
                    'isAdditional': False,
                    'weekType': parsed['weekType'],
                    'weekStart': parsed['weekStart'],
                    'weekEnd': parsed['weekEnd'],
                    'time': time_slot,
                    'timeStart': start_slot,
                    'timeEnd': end_slot
                }
                if 'weeks' in parsed:
                    lesson_obj['weeks'] = parsed['weeks']
                if 'customTime' in parsed:
                    lesson_obj['customTime'] = parsed['customTime']

                days_map[meta['dayIdx']]['lessons'].append(lesson_obj)

        days_list = []
        for dn, di in DAYS_ORDER:
            d = days_map[di]
            if d['lessons']:
                days_list.append({
                    'dayName': d['name'],
                    'dayIndex': d['idx'],
                    'lessons': d['lessons']
                })

        result_groups.append({
            'id': g_info['group'],
            'group': g_info['group'],
            'course': g_info['course'],
            'major': g_info['major'],
            'days': days_list,
            'shortGroup': g_info['shortGroup'],
            'specialization': g_info['specialization']
        })

    return result_groups


def preserve_ifmib_lesson_ids(database, previous):
    """Keep IDs for unchanged lessons when newly parsed rows alter order."""
    def entries(data):
        for group in data.get("groups", []):
            for day in group.get("days", []):
                for lesson in day.get("lessons", []):
                    yield group, day, lesson

    def key(group, day, lesson):
        return (group.get("id", group.get("group")), day.get("dayIndex"),
                lesson.get("timeStart"), lesson.get("timeEnd"), lesson.get("subject"),
                lesson.get("teacher"), lesson.get("room"))

    old_ids = defaultdict(deque)
    used = set()
    max_number = 0
    for group, day, lesson in entries(previous):
        identifier = str(lesson.get("id", ""))
        if identifier:
            old_ids[key(group, day, lesson)].append(identifier)
            used.add(identifier)
            match = re.search(r'-(\d+)$', identifier)
            if match:
                max_number = max(max_number, int(match.group(1)))

    for group, day, lesson in entries(database):
        candidates = old_ids[key(group, day, lesson)]
        if candidates:
            lesson["id"] = candidates.popleft()
        else:
            while True:
                max_number += 1
                identifier = f"{group.get('id', group.get('group', 'ifmib'))}-{max_number}"
                if identifier not in used:
                    used.add(identifier)
                    lesson["id"] = identifier
                    break


def save_ifmib_database(database, output_path):
    """Validate and atomically publish IFMiB data with sanity safeguards."""
    output_path = os.fspath(output_path)
    groups = database.get("groups", [])
    lessons = sum(len(day.get("lessons", [])) for group in groups for day in group.get("days", []))
    if not groups or not lessons:
        raise ValueError("IFMiB schedule is empty; previous output was preserved")

    if os.path.exists(output_path):
        with open(output_path, encoding="utf-8") as stream:
            previous = json.load(stream)
        preserve_ifmib_lesson_ids(database, previous)
        old_groups = previous.get("groups", [])
        old_lessons = sum(len(day.get("lessons", [])) for group in old_groups for day in group.get("days", []))
        if len(groups) < len(old_groups) * 0.7 or lessons < old_lessons * 0.7:
            raise ValueError(f"IFMiB schedule shrank: groups {len(old_groups)} -> {len(groups)}, lessons {old_lessons} -> {lessons}")

        if {k: v for k, v in previous.items() if k != "updatedAt"} == {k: v for k, v in database.items() if k != "updatedAt"}:
            database["updatedAt"] = previous.get("updatedAt", database.get("updatedAt"))
            print(f"✅ IFMiB schedule content is identical. Preserving updatedAt timestamp.")
    else:
        # Initial run: assign IDs
        counter = 0
        for group in groups:
            for day in group.get("days", []):
                for lesson in day.get("lessons", []):
                    counter += 1
                    lesson["id"] = f"{group['group']}-{counter}"

    directory = os.path.dirname(os.path.abspath(output_path))
    temporary_path = None
    try:
        with tempfile.NamedTemporaryFile(mode="w", encoding="utf-8", dir=directory, delete=False) as stream:
            temporary_path = stream.name
            json.dump(database, stream, ensure_ascii=False, indent=2)
            stream.flush()
            os.fsync(stream.fileno())
        os.replace(temporary_path, output_path)
    finally:
        if temporary_path and os.path.exists(temporary_path):
            os.unlink(temporary_path)

    print(f"Success! Saved {len(groups)} groups and {lessons} lessons to {output_path}")


def main():
    print("Finding IFMiB schedule URLs...")
    url_items = get_ifmib_xlsx_urls()
    if not url_items:
        print("Failed to find any IFMiB schedule URLs", file=sys.stderr)
        sys.exit(1)

    print(f"Found {len(url_items)} course files to parse:")
    for it in url_items:
        print(f" - {it['label']}: {it['url']}")

    all_groups = []
    import io
    for it in url_items:
        print(f"Downloading and parsing {it['label']}...")
        content = None
        for attempt in range(3):
            try:
                if requests is not None:
                    resp = requests.get(it['url'], headers=HEADERS, timeout=45)
                    resp.raise_for_status()
                    content = resp.content
                else:
                    import urllib.request
                    req = urllib.request.Request(it['url'], headers=HEADERS)
                    with urllib.request.urlopen(req, timeout=45) as resp:
                        content = resp.read()
                if content:
                    break
            except Exception as e:
                print(f"[Warning] Download attempt {attempt+1}/3 failed for {it['label']}: {e}", file=sys.stderr)
                if attempt < 2:
                    import time
                    time.sleep(2)

        if not content:
            print(f"Failed to download {it['label']}", file=sys.stderr)
            sys.exit(1)

        wb = openpyxl.load_workbook(io.BytesIO(content), data_only=True)
        course_groups = parse_schedule_sheet(wb, it['course'], it['isAngl'])
        print(f"   Parsed {len(course_groups)} subgroups from {it['label']}.")
        all_groups.extend(course_groups)

    all_groups.sort(key=lambda g: g["group"])

    iso_now = datetime.now(timezone.utc).strftime("%Y-%m-%dT%H:%M:%SZ")
    database = {
        "version": 4,
        "schemaVersion": "1.0",
        "semester": compute_semester_string(),
        "semesterMeta": compute_semester_meta(),
        "updatedAt": iso_now,
        "sourceUrl": IFMIB_SCHEDULE_URL,
        "groups": all_groups
    }

    base_dir = os.path.dirname(os.path.abspath(__file__))
    output_path = os.path.join(base_dir, OUTPUT_FILE)
    save_ifmib_database(database, output_path)


if __name__ == "__main__":
    main()
